"""Módulo de exportación de vacantes analizadas hacia Microsoft Excel.

Aplica estilos profesionales, auto-ajuste de columnas y filtros interactivos.
"""

import json
import logging
from typing import Optional
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.config import settings
from src.database import get_db_connection

logger = logging.getLogger(__name__)


def export_jobs_to_excel(
    output_path: Optional[str] = None, only_qualified: bool = False
) -> str:
    """Extrae las vacantes de SQLite y genera el informe formateado en Excel."""
    out_file = output_path or str(settings.EXCEL_REPORT_PATH)
    settings.OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    conn = get_db_connection()
    cursor = conn.cursor()

    query = "SELECT * FROM job_applications"
    if only_qualified:
        query += " WHERE status IN ('QUALIFIED', 'COMPILED')"
    query += " ORDER BY rowid DESC"

    cursor.execute(query)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospectos Calificados"

    header_fill = PatternFill(
        start_color="1A365D", end_color="1A365D", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    headers = [
        "Identificador",
        "Título del Rol",
        "Empresa",
        "Ubicación",
        "Score Total",
        "Hard Match (%)",
        "Soft Match (%)",
        "Estado",
        "Justificación Técnica",
        "Enlace Vacante",
        "Ruta CV Word",
        "Ruta CV PDF",
        "Fecha",
    ]

    ws.append(headers)

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    for r in rows:
        row_dict = dict(r)
        
        # Mapeo tolerante de nombres alternativos de columna
        job_identifier = str(row_dict.get("id") or row_dict.get("job_id") or "")[:8]
        score = row_dict.get("match_score", 0.0) or 0.0
        hard = row_dict.get("hard_match_score", 0.0) or 0.0
        soft = row_dict.get("soft_match_score", 0.0) or 0.0

        row_data = [
            job_identifier,
            row_dict.get("title", "") or "",
            row_dict.get("company", "") or "",
            row_dict.get("location", "") or "",
            round(float(score), 1),
            round(float(hard), 1),
            round(float(soft), 1),
            row_dict.get("status", "PENDING") or "PENDING",
            row_dict.get("score_rationale", "") or "",
            row_dict.get("url", "") or "",
            row_dict.get("cv_docx_path", "") or "",
            row_dict.get("cv_pdf_path", "") or "",
            str(row_dict.get("scraped_at", ""))[:10],
        ]
        ws.append(row_data)

    # Formato de celdas y auto-ancho
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = data_font
            cell.border = thin_border
            if cell.column in [1, 5, 6, 7, 8, 13]:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            if cell.column == 8:
                val = str(cell.value)
                if val in ("QUALIFIED", "COMPILED"):
                    cell.fill = PatternFill(
                        start_color="C6F6D5", end_color="C6F6D5", fill_type="solid"
                    )
                elif val == "DISQUALIFIED":
                    cell.fill = PatternFill(
                        start_color="FED7D7", end_color="FED7D7", fill_type="solid"
                    )

    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

    wb.save(out_file)
    logger.info(f"Reporte Excel generado exitosamente en: {out_file}")
    return out_file


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    print("--- TEST EXPORTADOR EXCEL ---")
    generated_file = export_jobs_to_excel()
    print("Excel guardado en:", generated_file)